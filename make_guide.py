# -*- coding: utf-8 -*-
"""iPad(ネット無し)運用の手順書 (Word) を生成する"""
import os
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

here = os.path.dirname(os.path.abspath(__file__))
doc = Document()
st = doc.styles['Normal']
st.font.name = 'Yu Gothic'
st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Yu Gothic')
st.font.size = Pt(10.5)

def h(t, lv=1): doc.add_heading(t, lv)
def p(t): doc.add_paragraph(t)
def n(t): doc.add_paragraph(t, style='List Number')
def b(t): doc.add_paragraph(t, style='List Bullet')

doc.add_heading('サッカー公式記録アプリ iPad(ネット無し)運用手順', 0)
p('目的：会場にネットが無くても，iPadでアプリを開いて記録できるようにする．アプリは一度だけネット経由で「ホーム画面に追加」し，以降は機内モードでも起動する．')

h('0. 仕組みと注意')
b('アプリ本体(空の状態)をGitHub Pagesに置く．学校名・選手名などのデータは置かない．')
b('iPadでそのURLを開いて「ホーム画面に追加」すると，アプリ本体がiPadに保存され，ネット無しでも開ける．')
b('学校データ(選手名簿)は，各校の参加申込書(Excel)をiPad上で読み込んで登録する．公開サイトには載らない．')
b('記録データはiPad内に自動保存される．大会後にExcel出力して他の場所にも控えを残す．')
b('公開URLはネットを知っている人なら誰でも開ける(中身は空のアプリのみ)．学校名・大会名をURLやリポジトリ名に入れない．')

h('1. 公開(パソコンで1回だけ)')
n('github.com にサインインする．右上の「＋」→「New repository」．')
n('Repository name：soccer-record(学校名・大会名を入れない)．Public を選ぶ．「Add a README file」はチェックしない．「Create repository」．')
n('作成後の画面で「uploading an existing file」を押す．')
n('フォルダ dev\\pwa の中の次の6ファイルをドラッグして入れる：index.html, sw.js, manifest.json, icon-180.png, icon-192.png, icon-512.png．(build_pwa.py と make_guide.py は入れない)')
n('下の「Commit changes」を押す．')
n('Settings → Pages → Build and deployment の Source を「Deploy from a branch」，Branch を「main」「/(root)」にして Save．')
n('1〜2分後，同じ画面に公開URL(https://nilaykanay.github.io/soccer-record/)が表示される．')

h('2. iPadの準備(ネットがある場所で1回だけ)')
n('iPadのSafariで公開URLを開く．アプリの画面(新規／保存済み一覧…)が出ることを確認．')
n('下部または上部の共有ボタン(□から↑)→「ホーム画面に追加」→「追加」．')
n('Safariを閉じ，ホーム画面にできた「サッカー記録」アイコンから起動する(この起動方法で使う)．')
n('起動したまま10秒ほど待つ(アプリ本体の保存が完了する)．')

h('3. 事前テスト(大会前に必ず実施)')
n('iPadを機内モードにし，Wi-FiもOFFにする．')
n('ホーム画面のアイコンから起動し，画面が出ることを確認．')
n('「学校データ」タブで各校の参加申込書(Excel)を読み込む(ファイルアプリ/AirDrop/Teams等で事前にiPadに入れておく)．')
n('試験的に数件入力し，「保存」→アプリを完全に閉じる→再度起動して，入力内容が残っていることを確認．')
n('「試合記録データを保存（Excel）」でExcelが出力でき，ファイルアプリで開けることを確認．')
n('PDF出力が使えることを確認．')

h('4. 大会当日')
b('ホーム画面のアイコンから起動する(Safariのタブから開かない)．')
b('試合ごとに「保存」し，終わったらExcelも出力して別の場所(Teams/メール/AirDrop)へ控えを送る(ネットがある時に)．')
b('途中でiPadを再起動しても，ホーム画面から起動すれば内容は残る．ただし「Safariの履歴とWebサイトデータを消去」を行うと消える．')

h('5. アプリを更新したとき')
b('更新版のファイルをGitHubのリポジトリに上書きアップロードする．iPadはネットにつながった状態で2回起動すると新しい版になる(1回目で取得，2回目から反映)．')
b('大会直前の更新は避ける．更新したら，事前テスト(3)をもう一度行う．')

h('6. うまく動かないとき')
b('「ホーム画面に追加」が出ない：Safariで開いているか確認(他のアプリ内ブラウザでは出ない)．')
b('機内モードで真っ白：ネットのある状態で一度ホーム画面アイコンから起動し，10秒待ってから再試験する．')
b('それでも不可の場合：パソコンで記録する，または紙の公式記録を併用する．')

doc.save(os.path.join(here, 'iPad運用手順書_ネット無し.docx'))
print('saved')
